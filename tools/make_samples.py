"""Regenerate the fixtures under samples/ deterministically.

Run with: python tools/make_samples.py
Requires requirements-dev.txt (fpdf2) and requirements.txt (python-docx)
to be installed.
"""

import pathlib

from docx import Document
from fpdf import FPDF

SAMPLES_DIR = pathlib.Path(__file__).resolve().parent.parent / "samples"

TXT_MARKER = "The quick brown fox jumps over the lazy dog."
PDF_MARKER = "Aloud reads PDF documents aloud."
DOCX_MARKER = "Aloud reads Word documents aloud."

TXT_PARAGRAPHS = [
    "Aloud is a small offline text-to-speech reader built for PDF, TXT, "
    "and DOCX documents. It extracts text, splits it into chunks, and "
    "synthesizes speech using a local neural voice model.",
    TXT_MARKER,
    "This sample file exists to exercise the plain-text extraction path: "
    "encoding detection, paragraph splitting, and downstream synthesis. "
    "It is intentionally short so tests run quickly.",
    "The final paragraph closes out the fixture with a bit more filler "
    "text so the document has a realistic multi-paragraph shape, rather "
    "than being a single line repeated for padding.",
]


def make_sample_txt():
    path = SAMPLES_DIR / "sample.txt"
    path.write_text("\n\n".join(TXT_PARAGRAPHS) + "\n", encoding="utf-8")
    return path


def make_sample_long_txt():
    # Must be >= 5x the character count of sample.txt.
    base = "\n\n".join(TXT_PARAGRAPHS) + "\n"
    long_paragraphs = list(TXT_PARAGRAPHS)
    filler = (
        "This paragraph is repeated several times to pad sample_long.txt "
        "well past five times the length of sample.txt, so duration-scales-"
        "with-input-length tests have clear headroom. "
    )
    while len("\n\n".join(long_paragraphs)) < len(base) * 6:
        long_paragraphs.append(filler * 4)
    path = SAMPLES_DIR / "sample_long.txt"
    path.write_text("\n\n".join(long_paragraphs) + "\n", encoding="utf-8")
    return path


class SamplePDF(FPDF):
    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 10, f"Page {self.page_no()} of 2", align="C")


def make_sample_pdf():
    pdf = SamplePDF()
    pdf.set_auto_page_break(auto=True, margin=20)

    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, "Aloud Sample PDF")
    pdf.ln(4)
    pdf.multi_cell(
        0,
        8,
        "This is a two-page, text-based PDF used to exercise Aloud's PDF "
        "extraction path, including stripping the 'Page X of 2' footer "
        "artifact from the extracted body text.",
    )
    pdf.ln(4)
    pdf.multi_cell(0, 8, PDF_MARKER)

    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(
        0,
        8,
        "This second page contains more body text so the extractor has "
        "content to read across a page boundary, while the footer on "
        "both pages should be stripped rather than read aloud.",
    )

    path = SAMPLES_DIR / "sample.pdf"
    pdf.output(str(path))
    return path


def make_sample_docx():
    doc = Document()
    doc.add_heading("Aloud Sample Document", level=1)
    doc.add_paragraph(
        "This DOCX fixture exercises Aloud's Word-document extraction "
        "path, including distinguishing headings from body paragraphs."
    )
    doc.add_paragraph(DOCX_MARKER)
    doc.add_paragraph(
        "A final paragraph closes out the document with additional "
        "filler content for a realistic multi-paragraph structure."
    )
    path = SAMPLES_DIR / "sample.docx"
    doc.save(str(path))
    return path


def make_corrupt_pdf():
    path = SAMPLES_DIR / "corrupt.pdf"
    path.write_bytes(
        b"This is not a valid PDF file. It has a .pdf extension but its "
        b"bytes do not form a PDF document structure at all, on purpose, "
        b"for the error-path test.\n"
    )
    return path


def main():
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    made = [
        make_sample_txt(),
        make_sample_long_txt(),
        make_sample_pdf(),
        make_sample_docx(),
        make_corrupt_pdf(),
    ]
    for path in made:
        print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
