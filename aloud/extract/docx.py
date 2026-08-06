"""DOCX extraction via python-docx, distinguishing headings from paragraphs."""

from docx import Document as DocxDocument

from aloud.errors import EmptyDocumentError, ExtractionError
from aloud.extract.base import Block, Document


def extract(path):
    path = str(path)
    try:
        docx_doc = DocxDocument(path)
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX file: {path}") from exc

    blocks = []
    title = None
    for paragraph in docx_doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        is_heading = style_name.startswith("Heading") or style_name.startswith("Title")
        kind = "heading" if is_heading else "paragraph"
        if is_heading and title is None:
            title = text
        blocks.append(Block(kind=kind, text=text))

    if not blocks:
        raise EmptyDocumentError(f"No text content found in {path}")

    return Document(title=title, blocks=blocks, source_path=path)
